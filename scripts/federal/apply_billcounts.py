import csv, re
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def load(f):
    with open(f, newline='', encoding='utf-8') as fh:
        return list(csv.DictReader(fh))

# --- billcount lookup: (Name,State,Chamber) -> TotalSponsoredItems / BioguideId ---
bc_rows = load('billcount_FINAL.csv')
bc_map = {}
bg_map = {}
for r in bc_rows:
    key = (r['Name'], r['State'], r['Chamber'])
    val = r['TotalSponsoredItems'].strip()
    bc_map[key] = val if val else None
    bg_map[key] = r['BioguideId']

def bill_label(name, state, chamber):
    v = bc_map.get((name, state, chamber))
    if v:
        return f"{v} (total sponsored items: bills+resolutions+amendments combined, via api.congress.gov)"
    return "N/A - no bioguide ID match found in lookup"

# --- Load Senate data ---
SEN_FIELDS = ["Name","State","StateLegChamber","StateLegYears","StateLegTerms",
              "SenateYears","SenateCongresses","SenateTerms",
              "BillsIntroducedSenate","BillsIntroducedStateLeg","Notes/Sources"]
sen_rows = []
for f in ['data/source/senate/senators_A.csv',
          'data/source/senate/senators_B.csv',
          'data/source/senate/senators_C.csv',
          'data/source/senate/senate_batch2_completeness_gap.csv',
          'data/source/senate/senate_batch3_completeness_gap2.csv']:
    sen_rows.extend(load(f))
for r in sen_rows:
    r['BillsIntroducedSenate'] = bill_label(r['Name'], r['State'], 'Senate')
    r['BioguideId'] = bg_map.get((r['Name'], r['State'], 'Senate'), '')

# --- Load House data (repair overflow commas) ---
HOUSE_FIELDS = ['Name','State','District','StateLegChamber','StateLegYears','StateLegTerms',
                'HouseYears','HouseCongresses','HouseTerms','BillsIntroducedHouse',
                'BillsIntroducedStateLeg','Notes/Sources']
NCOL = len(HOUSE_FIELDS)

def load_repair(f):
    rows = []
    with open(f, newline='', encoding='utf-8') as fh:
        reader = csv.reader(fh)
        header = next(reader)
        for raw in reader:
            if len(raw) < NCOL:
                raw = raw + ['']*(NCOL-len(raw))
            elif len(raw) > NCOL:
                head = raw[:NCOL-1]
                tail = ','.join(raw[NCOL-1:])
                raw = head + [tail]
            rows.append(dict(zip(HOUSE_FIELDS, raw)))
    return rows

STATE_FIX = {'OH':'Ohio','PA':'Pennsylvania','RI':'Rhode Island','NC':'North Carolina','NY':'New York','OR':'Oregon'}

house_files_regular = [
 'house_batch1_group1.csv','house_batch1_california.csv','house_batch1_group3.csv',
 'house_batch2_groupA.csv','house_batch2_groupB.csv','house_batch2_groupC.csv',
 'house_batch3_groupA.csv','house_batch3_groupB.csv','house_batch3_groupC.csv',
 'house_batch4_FINAL.csv',
 'house_batch5_texas.csv','house_batch5_groupB.csv','house_batch5_groupC.csv',
 'house_batch6_completeness_gap.csv','house_batch7_completeness_gap2.csv',
]
house_rows = []
for f in house_files_regular:
    rows = load_repair('data/source/house/'+f)
    for r in rows:
        if r['State'] in STATE_FIX:
            r['State'] = STATE_FIX[r['State']]
    house_rows.extend(rows)
for r in house_rows:
    r['BillsIntroducedHouse'] = bill_label(r['Name'], r['State'], 'House')
    r['BioguideId'] = bg_map.get((r['Name'], r['State'], 'House'), '')

print('Senate:', len(sen_rows), 'House:', len(house_rows))

sen_by_key = {}
for r in sen_rows:
    sen_by_key.setdefault((r['Name'], r['State']), []).append(r)
house_by_key = {}
for r in house_rows:
    house_by_key.setdefault((r['Name'], r['State']), []).append(r)

# Both-chamber detection uses BioguideId (a stable per-person ID), not Name+State --
# three people (Chuck/Charles Schumer, Bob/Robert Menendez, Jim/James Inhofe) have
# different name spellings between their House-era and Senate-era source research,
# which a Name+State match alone would miss (and did, earlier in this project).
sen_by_bg = {}
for r in sen_rows:
    if r['BioguideId']:
        sen_by_bg.setdefault(r['BioguideId'], []).append(r)
house_by_bg = {}
for r in house_rows:
    if r['BioguideId']:
        house_by_bg.setdefault(r['BioguideId'], []).append(r)
both_bg = sorted(set(sen_by_bg) & set(house_by_bg))
both_keys = sorted(set(sen_by_key) & set(house_by_key))  # kept for reference/back-compat, unused below
print('Both-chamber individuals:', len(both_bg))

def start_year(s):
    m = re.search(r'(1[89]\d{2}|20\d{2})', s or '')
    return int(m.group(1)) if m else 9999

wb = Workbook()
FONT_NAME = "Arial"
header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
body_font = Font(name=FONT_NAME, size=10)
na_font = Font(name=FONT_NAME, size=10, italic=True, color="808080")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap_top = Alignment(wrap_text=True, vertical="top")

def style_header(ws, ncols, height=32):
    for c in range(1, ncols+1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
    ws.row_dimensions[1].height = height

def style_body(ws, last_row, ncols):
    for row in ws.iter_rows(min_row=2, max_row=last_row, min_col=1, max_col=ncols):
        for cell in row:
            is_na = isinstance(cell.value, str) and cell.value.strip().startswith("N/A")
            cell.font = na_font if is_na else body_font
            cell.alignment = wrap_top
            cell.border = border

# ---------- Sheet 0: Overview ----------
ws0 = wb.active
ws0.title = "Overview"
ws0.column_dimensions['A'].width = 50
ws0.column_dimensions['B'].width = 16
filled_count = sum(1 for v in bc_map.values() if v)
rows_overview = [
    ("Metric", "Count"),
    ("U.S. Senators (1979-present) with state legislature background", len(sen_rows)),
    ("U.S. House Members (1979-present) with state legislature background", len(house_rows)),
    ("Combined total rows (Senators + House Members tabs)", len(sen_rows)+len(house_rows)),
    ("Individuals who served in BOTH chambers of Congress (plus a state legislature)", len(both_bg)),
    ("States covered", 50),
    ("Rows with a verified total sponsored-legislation count (via api.congress.gov)", filled_count),
    ("Rows without a bioguide ID match (bill count left N/A)", (len(sen_rows)+len(house_rows)) - filled_count),
]
for r in rows_overview:
    ws0.append(r)
style_header(ws0, 2, height=20)
for row in ws0.iter_rows(min_row=2, max_row=len(rows_overview), min_col=1, max_col=2):
    for cell in row:
        cell.font = body_font
        cell.border = border
ws0.append([])
ws0.append(["See tab \"Both Chambers\" for the 49 individuals who moved state legislature → one chamber of Congress → the other.", ""])
ws0.cell(row=ws0.max_row, column=1).font = Font(name=FONT_NAME, italic=True, size=10)
ws0.append(["Bill counts are TOTAL SPONSORED ITEMS (bills + resolutions + amendments combined), not bills alone — see Methodology tab for why a pure bills-only breakdown wasn't reliable to produce.", ""])
ws0.cell(row=ws0.max_row, column=1).font = Font(name=FONT_NAME, italic=True, size=10)

# ---------- Sheet: Both Chambers ----------
ws_both = wb.create_sheet("Both Chambers")
both_headers = ["Name","State","State Legislature Chamber(s)","State Legislature Years","State Legislature Terms",
                "U.S. House Years","U.S. House Congresses","U.S. House Terms",
                "U.S. Senate Years","U.S. Senate Congresses","U.S. Senate Terms",
                "Total Sponsored Items – House","Total Sponsored Items – Senate","Bills Introduced – State Legislature",
                "Sources / Notes"]
ws_both.append(both_headers)
for bg in both_bg:
    hrow = house_by_bg[bg][0]
    srow = sen_by_bg[bg][0]
    # Prefer the Senate-era name/state as canonical when the two records used different spellings
    name = srow.get('Name') or hrow.get('Name')
    state = srow.get('State') or hrow.get('State')
    stateleg_chamber = hrow.get('StateLegChamber') or srow.get('StateLegChamber')
    stateleg_years = hrow.get('StateLegYears') or srow.get('StateLegYears')
    stateleg_terms = hrow.get('StateLegTerms') or srow.get('StateLegTerms')
    name_note = "" if hrow.get('Name') == srow.get('Name') else f" [House record used name \"{hrow.get('Name')}\"]"
    notes = f"House record: {hrow.get('Notes/Sources','')} || Senate record: {srow.get('Notes/Sources','')}{name_note}"
    ws_both.append([
        name, state, stateleg_chamber, stateleg_years, stateleg_terms,
        hrow.get('HouseYears'), hrow.get('HouseCongresses'), hrow.get('HouseTerms'),
        srow.get('SenateYears'), srow.get('SenateCongresses'), srow.get('SenateTerms'),
        hrow.get('BillsIntroducedHouse'), srow.get('BillsIntroducedSenate'), hrow.get('BillsIntroducedStateLeg'),
        notes
    ])
style_header(ws_both, len(both_headers))
last = ws_both.max_row
style_body(ws_both, last, len(both_headers))
widths_both = [18,14,26,18,14,18,16,14,18,16,14,30,30,20,50]
for i,w in enumerate(widths_both, start=1):
    ws_both.column_dimensions[get_column_letter(i)].width = w
ws_both.freeze_panes = "A2"
ws_both.auto_filter.ref = f"A1:{get_column_letter(len(both_headers))}{last}"

# ---------- Sheet: Senators ----------
ws_sen = wb.create_sheet("Senators")
sen_headers = ["Name", "State", "State Legislature Chamber(s)", "State Legislature Years",
    "State Legislature Terms", "U.S. Senate Years", "U.S. Senate Congresses",
    "U.S. Senate Terms", "Total Sponsored Items – U.S. Senate",
    "Bills Introduced – State Legislature", "Sources / Notes"]
ws_sen.append(sen_headers)
sen_sorted = sorted(sen_rows, key=lambda r: (r['State'], start_year(r['SenateYears'])))
for r in sen_sorted:
    ws_sen.append([r[k] for k in SEN_FIELDS])
style_header(ws_sen, len(sen_headers))
last = ws_sen.max_row
style_body(ws_sen, last, len(sen_headers))
widths_sen = [20, 14, 26, 20, 16, 20, 16, 22, 36, 22, 46]
for i,w in enumerate(widths_sen, start=1):
    ws_sen.column_dimensions[get_column_letter(i)].width = w
ws_sen.freeze_panes = "A2"
ws_sen.auto_filter.ref = f"A1:{get_column_letter(len(sen_headers))}{last}"

# ---------- Sheet: House Members ----------
ws_house = wb.create_sheet("House Members")
house_headers = [
    "Name", "State", "District(s)", "State Legislature Chamber(s)", "State Legislature Years",
    "State Legislature Terms", "U.S. House Years", "U.S. House Congresses",
    "U.S. House Terms", "Total Sponsored Items – U.S. House",
    "Bills Introduced – State Legislature", "Sources / Notes"]
ws_house.append(house_headers)
house_sorted = sorted(house_rows, key=lambda r: (r['State'], start_year(r['HouseYears'])))
for r in house_sorted:
    ws_house.append([r.get(k,'') for k in HOUSE_FIELDS])
style_header(ws_house, len(house_headers))
last = ws_house.max_row
style_body(ws_house, last, len(house_headers))
widths_house = [20, 16, 14, 30, 22, 18, 20, 18, 24, 36, 22, 50]
for i,w in enumerate(widths_house, start=1):
    ws_house.column_dimensions[get_column_letter(i)].width = w
ws_house.freeze_panes = "A2"
ws_house.auto_filter.ref = f"A1:{get_column_letter(len(house_headers))}{last}"

# ---------- Sheet: Methodology ----------
ws_m = wb.create_sheet("Methodology & Sources")
ws_m.column_dimensions['A'].width = 112
notes = [
    ("U.S. Legislators Who Served in Both a State Legislature and Congress (1979–present)", True, 13),
    ("", False, 10),
    ("Status: COMPLETE, including bill-activity counts for all 959 rows (100%)", True, 12),
    (f"{len(sen_rows)} Senators and {len(house_rows)} House members (959 rows total) who served in Congress "
     "at any point from January 3, 1979 through August 2026, and who also held a seat in a state house/"
     "assembly or state senate at some point in their career — not necessarily in the same state they "
     "later represented in Congress.", False, 10),
    ("", False, 10),
    ("Bill-count methodology — IMPORTANT, read before citing these numbers", True, 11),
    ("The “Total Sponsored Items” columns come from the official api.congress.gov API (Library of "
     "Congress), obtained with a free API key the user registered. Each number is the API's own top-level "
     "“count” field for that member's /sponsored-legislation endpoint — the TOTAL of bills, "
     "resolutions, AND floor amendments they sponsored across their entire career, all combined into one "
     "figure. This is NOT a pure “bills introduced” count.", False, 10),
    ("A pilot test attempted to split this into separate bills/resolutions/amendments counts by paginating "
     "through every individual item per person and classifying it by type. That pilot failed validation: "
     "only 45% of a 20-person test sample reconciled cleanly against the API's own official total, with "
     "failure modes including pages that falsely claimed to be complete, a page that silently dropped "
     "~68 items with no error, and one person where two fetches of identical data produced contradictory "
     "results. The root cause is that this session has no raw HTTP/JSON access — API responses are read "
     "through a tool that summarizes them via a smaller AI model rather than parsing exact structured data, "
     "which is reliable for a single number (the combined total) but not for enumerating and classifying "
     "hundreds of individual records per person. Given that finding, the combined-total approach was used "
     "instead, as the more trustworthy option, at the user's direction.", False, 10),
    ("", False, 10),
    ("Bioguide ID matching — a two-pass process", True, 11),
    ("The API has no name-search endpoint, so a lookup table (2,036 unique members) was built from the "
     "per-Congress member-roster endpoint (96th–119th Congresses) and matched to this project's 959 people "
     "by last name + state, with first-name disambiguation where a state had multiple same-surname members "
     "(e.g. father/son pairs like Ray LaHood & Darin LaHood, or unrelated members who share a surname). "
     "This caught and corrected 18 initial name-collision errors (~35 individual people) where a naive "
     "match would have silently attributed one person's sponsorship count to a different person of the same "
     "surname — each was manually verified and corrected via independent search "
     "(e.g. Phillip Burton vs. his brother John Burton, both California Representatives, previously "
     "conflated; Carl D. Perkins vs. his son Carl C. Perkins, both Kentucky Representatives). "
     "The remaining 96 of 959 people (10%) had no match in that lookup table (mostly older, shorter-tenured "
     "members from the 1979–1994 era where the lookup build hit some incomplete API pages) — these were "
     "resolved in a second pass via individual targeted search (bioguide ID confirmed against congress.gov, "
     "history.house.gov, GovTrack, or Wikipedia for each), again with explicit attention to father/son and "
     "same-surname collisions (e.g. Harold Ford Sr. vs. Jr., both Tennessee Representatives; Harley O. "
     "Staggers Jr. vs. Sr., both West Virginia). All 96 were resolved with high confidence and none were "
     "left unresolved, giving 100% bill-count coverage across the full dataset.", False, 10),
    ("", False, 10),
    ("State-legislature bill counts", True, 11),
    ("“Bills Introduced – State Legislature” remains N/A for most rows, as documented previously: "
     "most state legislatures did not digitize bill-introduction records until roughly the 1990s–2000s, "
     "and there is no unified cross-state historical database comparable to Congress.gov.", False, 10),
    ("", False, 10),
    ("This workbook's tabs", True, 11),
    ("Overview: summary counts. Both Chambers: the 49 individuals who served in a state legislature, the "
     "U.S. House, AND the U.S. Senate. Senators: the full 136-row Senate table. House Members: the full "
     "823-row House table.", False, 10),
    ("", False, 10),
    ("Companion dataset: full bill-by-bill data with topics (Stata files, not in this workbook)", True, 11),
    ("At the user's request, every individual sponsored bill, resolution, and floor amendment for all 910 "
     "unique people in this project (not capped, not sampled) was separately extracted via a plain Python "
     "script run directly against the api.congress.gov API by the user on their own computer — deliberately "
     "outside this session's tools, because this session's web-fetch tool summarizes page content through a "
     "smaller AI model rather than parsing exact JSON, which is unsuitable at this scale. That script parses "
     "raw JSON with no model in the loop and self-checks every person's collected count against the API's "
     "own reported total. Two companion Stata files were produced: person_level.dta (910 rows, one per "
     "legislator, merge key bioguide_id) and sponsored_legislation_full.dta (202,617 rows, one per "
     "sponsored item, with congress, type, introduced date, title, Congress.gov Policy Area classification, "
     "and latest action — merges m:1 onto person_level.dta via bioguide_id).", False, 10),
    ("Two data-quality findings from that extraction, corrected before finalizing:", False, 10),
    ("(1) Three people — Chuck Schumer, Bob Menendez, and Jim Inhofe — had been treated throughout this "
     "project as two separate people each, because their House-era and Senate-era source research used "
     "different name spellings (e.g. \"Chuck Schumer\" vs. \"Charles Schumer\") that a name+state match "
     "didn't catch. Re-keying everything on bioguide_id (a stable per-person ID, not a name string) fixed "
     "this: the correct both-chamber count is 49, not the 46 reported earlier in this project, and each of "
     "those three now has one unified record instead of two duplicated ones.", False, 10),
    ("(2) For 33 of the 910 people — all high-volume sponsors needing more than 250 lifetime items — "
     "Congress.gov's own summary \"count\" statistic runs slightly ahead (0.03%-1.4%, max 45 of 3,329 items "
     "for Howard Metzenbaum) of what its itemized list endpoint actually returns, confirmed by re-fetching "
     "those 33 people a second time with extra pagination margin and finding byte-for-byte the same set of "
     "items both times, including cases where the API returned a genuinely empty page before the safety "
     "margin was even used up. This is a real characteristic of Congress.gov's own data, not a gap in this "
     "extraction — person_level.dta's total_sponsored_items column preserves the API's original stated "
     "total for reference, while sponsored_legislation_full.dta contains everything actually retrievable.", False, 10),
    ("Two further caveats specific to the bill-level file: floor amendments (about 26% of all items) "
     "structurally have no title or Policy Area in the API — Congress.gov does not assign topics to "
     "amendments, so this is a real data characteristic, not a missing-data error, and is exactly why the "
     "user's original question about distinguishing bills from amendments matters. Separately, about 1.9% "
     "of Bill/Resolution rows — all from the 93rd-95th Congresses (1973-1979), i.e. career activity that "
     "predates this project's 1979-present scope but is included because it's part of the same person's "
     "full sponsorship history — carry an older, more granular legacy subject-heading term instead of one "
     "of the ~34 standard modern Policy Area categories; these are left as returned by the API rather than "
     "remapped, since inventing a mapping wasn't asked for.", False, 10),
    ("", False, 10),
    ("bill_status column — how it was built and its limits", True, 11),
    ("Congress.gov does not provide a clean, pre-standardized status field (e.g. \"Passed\"/\"Failed\"/"
     "\"Became Law\") on the sponsored-legislation endpoint — only free-text latest_action_text, the single "
     "most recently recorded legislative action for that item. That text follows fairly formulaic phrasing "
     "generated by the same congressional record-keeping process, so — the same approach GovTrack.us takes "
     "over this same underlying Congress.gov data — it can be classified into a standardized status via "
     "rule-based pattern matching. This was done entirely on already-downloaded data (no new API calls): a "
     "set of ~40 regex rules, applied in priority order (most specific/conclusive status checked first — "
     "e.g. \"Became Law\" before \"Referred to Committee\"), was built from empirical frequency analysis of "
     "the actual latest_action_text phrasing across all 202,617 rows, then iteratively refined against "
     "samples of whatever didn't yet match.", False, 10),
    ("Because latest_action_text is a single most-recent snapshot rather than a full action history, "
     "bill_status reflects a bill's most recently known state, not necessarily every stage it passed "
     "through — a bill that Became Law will show \"Became Law\" even though it was also, earlier, "
     "\"Referred to Committee\" and \"Passed House/Senate\" along the way.", False, 10),
    ("Coverage: 99.1% of the 150,712 Bill/Resolution rows and 99.98% of the 51,905 Amendment rows matched "
     "a specific rule. The remainder is labeled \"Unclassified\" rather than guessed — mostly rare or "
     "idiosyncratic phrasings (one-off procedural notices, typos in the source record, cross-references to "
     "other bill numbers) that didn't reliably fit a general pattern. Bill/Resolution status categories: "
     "Referred to Committee (75.1%), Reported by Committee / On Calendar (9.7%), Other Procedural Action "
     "(6.4% — e.g. motions to reconsider, inter-chamber messages, cross-references — actions that don't by "
     "themselves indicate pass/fail), Became Law (3.5%), Agreed To [Resolution] (2.9%), Unclassified (0.9%), "
     "Failed/Rejected (0.6%), Held at the Desk (0.4%), Passed Senate (0.3%), Passed House (0.3%), Vetoed by "
     "President (0.02%), Vetoed [Veto Not Overridden] (0.01%). Amendment status categories use different, "
     "amendment-specific vocabulary: No Recorded Action (56.2% — Congress.gov often never logs a final "
     "action for floor amendments), Amendment Agreed To (30.5%), Amendment Not Agreed To/Failed (7.7%), "
     "Amendment Withdrawn (3.3%), Amendment Ruled Out of Order/Non-Germane (0.9%), Considered/Proposed with "
     "No Recorded Disposition (0.6%), Point of Order Raised with Outcome Not in the Text (0.5% — flagged "
     "separately rather than guessed, since the record doesn't say whether the point of order was sustained "
     "or overruled), Referred to Committee (0.3%), Unclassified (0.02%). The full classification script "
     "(classify_status.py) and per-category counts (status_classification_report.txt) are available on "
     "request for anyone who wants to audit or extend the rules.", False, 10),
    ("", False, 10),
    ("Compiled", True, 11),
    ("Research conducted August 2026 via Wikipedia, Ballotpedia, web search, and the api.congress.gov API. "
     "Prepared for the SDSU MA Political Science project.", False, 10),
]
r = 1
for text, bold, size in notes:
    cell = ws_m.cell(row=r, column=1, value=text)
    cell.font = Font(name=FONT_NAME, bold=bold, size=size)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

wb.save("US_Legislators_StateLegislature_and_Congress_1979_present.xlsx")
print("Saved combined workbook with bill counts.")
